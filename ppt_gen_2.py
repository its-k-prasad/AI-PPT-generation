
import streamlit as st
import json
import io
import requests
from typing import List, Dict, Any, Optional
from dataclasses import dataclass
from datetime import datetime
import base64
import tempfile
import os
import PyPDF2
import docx
import pandas as pd
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image
from reportlab.lib.colors import HexColor, black, white
from reportlab.pdfgen import canvas
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY
from reportlab.platypus.tableofcontents import TableOfContents
import google.generativeai as genai

# ============================================================================
# CONFIGURATION & SETUP
# ============================================================================

st.set_page_config(
    page_title="AI PDF Presentation Generator",
    page_icon="📊",
    layout="wide"
)

# ============================================================================
# DATA MODELS
# ============================================================================

@dataclass
class SlideContent:
    title: str
    bullet_points: List[str]
    additional_info: str = ""
    image_urls: List[str] = None
    reference_urls: List[str] = None

    def __post_init__(self):
        if self.image_urls is None:
            self.image_urls = []
        if self.reference_urls is None:
            self.reference_urls = []

@dataclass
class PresentationData:
    slides: List[SlideContent]
    topic: str
    generated_at: datetime

# ============================================================================
# DOCUMENT PROCESSING MODULE
# ============================================================================

class DocumentProcessor:
    """Handle various document formats and extract text content"""
    
    @staticmethod
    def extract_text_from_pdf(file) -> str:
        try:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            return text.strip()
        except Exception as e:
            st.error(f"Error reading PDF: {str(e)}")
            return ""

    @staticmethod  
    def extract_text_from_docx(file) -> str:
        try:
            doc = docx.Document(file)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            st.error(f"Error reading DOCX: {str(e)}")
            return ""

    @staticmethod
    def extract_text_from_txt(file) -> str:
        try:
            return str(file.read(), "utf-8")
        except Exception as e:
            st.error(f"Error reading TXT: {str(e)}")
            return ""

    @staticmethod
    def extract_text_from_csv(file) -> str:
        try:
            df = pd.read_csv(file)
            summary = f"Dataset Summary:\n"
            summary += f"Shape: {df.shape[0]} rows, {df.shape[1]} columns\n"
            summary += f"Columns: {', '.join(df.columns.tolist())}\n\n"
            summary += "Data Overview:\n"
            summary += df.describe(include='all').to_string()
            return summary
        except Exception as e:
            st.error(f"Error reading CSV: {str(e)}")
            return ""

    def process_uploaded_file(self, uploaded_file) -> str:
        if uploaded_file is None:
            return ""
            
        file_type = uploaded_file.type
        
        if file_type == "application/pdf":
            return self.extract_text_from_pdf(uploaded_file)
        elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self.extract_text_from_docx(uploaded_file)
        elif file_type == "text/plain":
            return self.extract_text_from_txt(uploaded_file)
        elif file_type == "text/csv":
            return self.extract_text_from_csv(uploaded_file)
        else:
            st.error(f"Unsupported file type: {file_type}")
            return ""

# ============================================================================
# AI CONTENT GENERATION MODULE
# ============================================================================

class AIContentGenerator:
    """Generate presentation content using Google Gemini"""
    
    def __init__(self, api_key: str):
        self.api_key = api_key
        if api_key:
            genai.configure(api_key=api_key)
        
    def create_presentation_prompt(self, topic: str, document_content: str = "") -> str:
        """Create the system prompt for enhanced presentation generation"""
        
        base_prompt = """You are an intelligent assistant that helps generate professional PDF presentations with rich content.

Your task is to generate comprehensive presentation-ready slide content with additional information, relevant image suggestions, and reference URLs.

Instructions:
1. Create 6-10 slides including title, introduction, content slides, and conclusion
2. Each slide should have a clear title and 3-6 concise bullet points
3. Include additional_info with detailed explanations, statistics, or examples
4. Suggest relevant image_urls (use placeholder URLs like https://via.placeholder.com/400x300/0066cc/ffffff?text=Topic+Image)
5. Provide reference_urls with credible sources and further reading links
6. Keep content professional, informative, and comprehensive

Output Format - Respond ONLY with valid JSON in this exact structure:
```json
{
  "slides": [
    {
      "title": "Slide Title",
      "bullet_points": [
        "First key point",
        "Second important point", 
        "Third valuable insight"
      ],
      "additional_info": "Detailed explanation, statistics, examples, or context that supports the main points. This should be informative and add value to the presentation.",
      "image_urls": [
        "https://via.placeholder.com/400x300/0066cc/ffffff?text=Relevant+Image1",
        "https://via.placeholder.com/400x300/28a745/ffffff?text=Chart+or+Graph"
      ],
      "reference_urls": [
        "https://example.com/relevant-article",
        "https://research.example.com/study"
      ]
    }
  ]
}
```
"""
        
        if document_content:
            content_preview = document_content[:3000] + "..." if len(document_content) > 3000 else document_content
            prompt = f"{base_prompt}\n\nDocument Content to Analyze:\n{content_preview}\n\nTopic: {topic}"
        else:
            prompt = f"{base_prompt}\n\nTopic to Present: {topic}\n\nGenerate comprehensive content based on your knowledge with relevant examples, statistics, and references."
            
        return prompt

    def generate_with_gemini(self, prompt: str) -> Optional[Dict]:
        """Generate content using Google Gemini"""
        try:
            model = genai.GenerativeModel('gemini-2.0-flash')
            response = model.generate_content(prompt)
            
            response_text = response.text.strip()
            
            if "```json" in response_text:
                json_start = response_text.find("```json") + 7
                json_end = response_text.rfind("```")
                if json_end > json_start:
                    response_text = response_text[json_start:json_end].strip()
            
            json_start = response_text.find('{')
            json_end = response_text.rfind('}') + 1
            
            if json_start >= 0 and json_end > json_start:
                json_str = response_text[json_start:json_end]
                return json.loads(json_str)
            else:
                st.error("Could not parse JSON from AI response")
                return None
                
        except Exception as e:
            st.error(f"Error generating content with Gemini: {str(e)}")
            return None

    def generate_presentation_content(self, topic: str, document_content: str = "") -> Optional[PresentationData]:
        """Main method to generate enhanced presentation content"""
        
        if not self.api_key:
            st.error("Please provide a valid Gemini API key")
            return self._generate_fallback_content(topic)
        
        prompt = self.create_presentation_prompt(topic, document_content)
        result = self.generate_with_gemini(prompt)
        
        if not result:
            return self._generate_fallback_content(topic)
        
        try:
            slides = []
            for slide_data in result.get('slides', []):
                slides.append(SlideContent(
                    title=slide_data.get('title', ''),
                    bullet_points=slide_data.get('bullet_points', []),
                    additional_info=slide_data.get('additional_info', ''),
                    image_urls=slide_data.get('image_urls', []),
                    reference_urls=slide_data.get('reference_urls', [])
                ))
            
            return PresentationData(
                slides=slides,
                topic=topic,
                generated_at=datetime.now()
            )
            
        except Exception as e:
            st.error(f"Error parsing AI response: {str(e)}")
            return self._generate_fallback_content(topic)

    def _generate_fallback_content(self, topic: str) -> PresentationData:
        """Generate enhanced fallback content when AI fails"""
        slides = [
            SlideContent(
                title=f"{topic}",
                bullet_points=["AI-Generated Professional Presentation", "Comprehensive Analysis", "Data-Driven Insights"],
                additional_info="This presentation has been automatically generated using advanced AI technology to provide comprehensive coverage of the topic with relevant insights and supporting information.",
                image_urls=["https://via.placeholder.com/400x300/1f497d/ffffff?text=Title+Slide"],
                reference_urls=["https://example.com/ai-presentations"]
            ),
            SlideContent(
                title="Introduction & Overview",
                bullet_points=[
                    f"Comprehensive overview of {topic}",
                    "Key concepts and fundamental principles",
                    "Current relevance and importance",
                    "Presentation structure and objectives"
                ],
                additional_info="This section provides the foundational understanding necessary to grasp the core concepts. We'll explore the historical context, current applications, and future implications of the topic.",
                image_urls=["https://via.placeholder.com/400x300/28a745/ffffff?text=Overview"],
                reference_urls=["https://example.com/introduction", "https://example.com/overview"]
            ),
            SlideContent(
                title="Key Analysis & Findings",
                bullet_points=[
                    "Primary research findings and data",
                    "Statistical analysis and trends",
                    "Critical insights and interpretations",
                    "Practical applications and use cases"
                ],
                additional_info="Our analysis reveals significant patterns and trends that have important implications for understanding this topic. These findings are based on current research and real-world applications.",
                image_urls=["https://via.placeholder.com/400x300/dc3545/ffffff?text=Data+Analysis"],
                reference_urls=["https://example.com/research", "https://example.com/data-analysis"]
            ),
            SlideContent(
                title="Conclusions & Next Steps",
                bullet_points=[
                    "Summary of critical findings",
                    "Key takeaways and insights",
                    "Future considerations and opportunities",
                    "Recommended actions and next steps"
                ],
                additional_info="Based on our comprehensive analysis, these conclusions provide a clear path forward. The recommendations are actionable and based on evidence-driven insights.",
                image_urls=["https://via.placeholder.com/400x300/6f42c1/ffffff?text=Conclusions"],
                reference_urls=["https://example.com/conclusions", "https://example.com/future-research"]
            )
        ]
        
        return PresentationData(
            slides=slides,
            topic=topic,
            generated_at=datetime.now()
        )

# ============================================================================
# ENHANCED PDF GENERATION MODULE
# ============================================================================

class EnhancedPDFGenerator:
    """Generate enhanced PDF presentations with images and rich content"""
    
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()
        
    def _setup_custom_styles(self):
        """Setup custom styles for enhanced PDF"""
        
        self.styles.add(ParagraphStyle(
            name='TitleSlide',
            parent=self.styles['Title'],
            fontSize=28,
            spaceAfter=30,
            alignment=TA_CENTER,
            textColor=HexColor('#1f497d'),
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='SlideTitle',
            parent=self.styles['Heading1'],
            fontSize=22,
            spaceAfter=15,
            textColor=HexColor('#1f497d'),
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='BulletPoint',
            parent=self.styles['Normal'],
            fontSize=12,
            spaceAfter=8,
            leftIndent=20,
            textColor=HexColor('#333333')
        ))
        
        self.styles.add(ParagraphStyle(
            name='AdditionalInfo',
            parent=self.styles['Normal'],
            fontSize=11,
            spaceAfter=10,
            alignment=TA_JUSTIFY,
            textColor=HexColor('#555555'),
            borderWidth=1,
            borderColor=HexColor('#e0e0e0'),
            borderPadding=10,
            backColor=HexColor('#f8f9fa')
        ))
        
        self.styles.add(ParagraphStyle(
            name='URLStyle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=HexColor('#0066cc'),
            leftIndent=10
        ))
        
    def download_image(self, url: str) -> Optional[str]:
        """Download image from URL and return local path"""
        try:
            response = requests.get(url, stream=True, timeout=10)
            response.raise_for_status()
            
            # Create temporary file
            with tempfile.NamedTemporaryFile(delete=False, suffix='.jpg') as tmp_file:
                for chunk in response.iter_content(chunk_size=8192):
                    tmp_file.write(chunk)
                return tmp_file.name
        except Exception as e:
            print(f"Error downloading image {url}: {e}")
            return None
    
    def create_presentation(self, presentation_data: PresentationData) -> io.BytesIO:
        """Create enhanced PDF presentation"""
        
        pdf_buffer = io.BytesIO()
        
        doc = SimpleDocTemplate(
            pdf_buffer,
            pagesize=A4,
            rightMargin=50,
            leftMargin=50,
            topMargin=60,
            bottomMargin=60
        )
        
        story = []
        
        # Generate enhanced slides
        for i, slide_content in enumerate(presentation_data.slides):
            if i == 0:
                story.extend(self._create_enhanced_title_slide(slide_content, presentation_data))
            else:
                story.extend(self._create_enhanced_content_slide(slide_content, i))
            
            if i < len(presentation_data.slides) - 1:
                story.append(PageBreak())
        
        # Build PDF
        doc.build(story)
        pdf_buffer.seek(0)
        
        return pdf_buffer

    def _create_enhanced_title_slide(self, slide_content: SlideContent, presentation_data: PresentationData) -> List:
        """Create enhanced title slide"""
        content = []
        
        content.append(Spacer(1, 1.5*inch))
        
        # Main title
        title = Paragraph(presentation_data.topic, self.styles['TitleSlide'])
        content.append(title)
        content.append(Spacer(1, 0.5*inch))
        
        # Subtitle points
        if slide_content.bullet_points:
            for point in slide_content.bullet_points:
                subtitle = Paragraph(f"• {point}", self.styles['BulletPoint'])
                content.append(subtitle)
        
        content.append(Spacer(1, 0.3*inch))
        
        # Additional info
        if slide_content.additional_info:
            info_para = Paragraph(slide_content.additional_info, self.styles['AdditionalInfo'])
            content.append(info_para)
        
        content.append(Spacer(1, 0.5*inch))
        
        # Add title image if available
        if slide_content.image_urls:
            image_path = self.download_image(slide_content.image_urls[0])
            if image_path:
                try:
                    img = Image(image_path, width=3*inch, height=2*inch)
                    content.append(img)
                    os.unlink(image_path)  # Clean up temp file
                except:
                    pass
        
        # Timestamp
        timestamp = f"Generated: {presentation_data.generated_at.strftime('%B %d, %Y')}"
        footer = Paragraph(timestamp, self.styles['URLStyle'])
        content.append(Spacer(1, 0.3*inch))
        content.append(footer)
        
        return content

    def _create_enhanced_content_slide(self, slide_content: SlideContent, slide_number: int) -> List:
        """Create enhanced content slide"""
        content = []
        
        content.append(Spacer(1, 0.3*inch))
        
        # Slide title
        title = Paragraph(slide_content.title, self.styles['SlideTitle'])
        content.append(title)
        content.append(Spacer(1, 0.2*inch))
        
        # Bullet points
        for bullet_point in slide_content.bullet_points:
            bullet_text = f"• {bullet_point}"
            bullet = Paragraph(bullet_text, self.styles['BulletPoint'])
            content.append(bullet)
        
        content.append(Spacer(1, 0.2*inch))
        
        # Additional information
        if slide_content.additional_info:
            info_para = Paragraph(f"<b>Additional Information:</b><br/>{slide_content.additional_info}", 
                                 self.styles['AdditionalInfo'])
            content.append(info_para)
            content.append(Spacer(1, 0.2*inch))
        
        # Add images
        if slide_content.image_urls:
            for img_url in slide_content.image_urls[:2]:  # Limit to 2 images per slide
                image_path = self.download_image(img_url)
                if image_path:
                    try:
                        img = Image(image_path, width=2.5*inch, height=1.8*inch)
                        content.append(img)
                        content.append(Spacer(1, 0.1*inch))
                        os.unlink(image_path)  # Clean up temp file
                    except:
                        pass
        
        # Reference URLs
        if slide_content.reference_urls:
            content.append(Spacer(1, 0.1*inch))
            ref_title = Paragraph("<b>References & Further Reading:</b>", self.styles['BulletPoint'])
            content.append(ref_title)
            
            for url in slide_content.reference_urls:
                url_para = Paragraph(f"• {url}", self.styles['URLStyle'])
                content.append(url_para)
        
        return content

# ============================================================================
# STREAMLIT APPLICATION
# ============================================================================

def main():
    """Main Streamlit application"""
    
    st.title("🚀 Enhanced AI PDF Presentation Generator")
    st.markdown("Generate professional PDF presentations with images, detailed information, and references")
    
    # API Key input
    api_key = st.text_input(
        "🔑 Gemini API Key",
        type="password",
        help="Get your API key from: https://makersuite.google.com/app/apikey"
    )
    
    if not api_key:
        st.warning("Please enter your Gemini API key to continue")
        return
    
    # Initialize components
    doc_processor = DocumentProcessor()
    ai_generator = AIContentGenerator(api_key)
    pdf_generator = EnhancedPDFGenerator()
    
    # Main interface
    st.header("📝 Create Your Presentation")
    
    # Input method selection
    input_method = st.radio(
        "Choose input method:",
        options=["Topic Only", "Upload Document"],
        horizontal=True
    )
    
    topic = ""
    document_content = ""
    
    if input_method == "Topic Only":
        topic = st.text_input(
            "Enter Presentation Topic",
            placeholder="e.g., Machine Learning in Healthcare"
        )
    else:
        topic = st.text_input(
            "Enter Presentation Topic",
            placeholder="e.g., Document Analysis Summary"
        )
        
        uploaded_file = st.file_uploader(
            "Upload Document",
            type=['pdf', 'docx', 'txt', 'csv']
        )
        
        if uploaded_file:
            with st.spinner("Processing document..."):
                document_content = doc_processor.process_uploaded_file(uploaded_file)
                
            if document_content:
                st.success(f"✅ Document processed! {len(document_content)} characters extracted")
    
    # Generate presentation
    if st.button("🎯 Generate Enhanced PDF Presentation", type="primary", disabled=not topic.strip()):
        with st.spinner("🤖 AI is creating your enhanced presentation..."):
            presentation_data = ai_generator.generate_presentation_content(topic, document_content)
            
            if presentation_data:
                st.session_state['presentation_data'] = presentation_data
                st.success("✅ Enhanced presentation generated!")
                st.balloons()
    
    # Display and download results
    if 'presentation_data' in st.session_state:
        presentation_data = st.session_state['presentation_data']
        
        st.header("📊 Generated Presentation")
        
        # Show slide previews
        for i, slide in enumerate(presentation_data.slides):
            with st.expander(f"Slide {i+1}: {slide.title}"):
                
                # Bullet points
                st.subheader("Key Points:")
                for bullet in slide.bullet_points:
                    st.write(f"• {bullet}")
                
                # Additional info
                if slide.additional_info:
                    st.subheader("Additional Information:")
                    st.info(slide.additional_info)
                
                # Images
                if slide.image_urls:
                    st.subheader("Related Images:")
                    cols = st.columns(len(slide.image_urls[:3]))
                    for idx, img_url in enumerate(slide.image_urls[:3]):
                        with cols[idx]:
                            st.image(img_url, caption=f"Image {idx+1}", width=200)
                
                # References
                if slide.reference_urls:
                    st.subheader("References:")
                    for url in slide.reference_urls:
                        st.write(f"🔗 {url}")
        
        # Download PDF
        st.header("📥 Download PDF")
        
        if st.button("📄 Generate PDF File", type="primary"):
            with st.spinner("Creating enhanced PDF..."):
                try:
                    pdf_buffer = pdf_generator.create_presentation(presentation_data)
                    
                    st.download_button(
                        label="⬇️ Download Enhanced PDF",
                        data=pdf_buffer.getvalue(),
                        file_name=f"{presentation_data.topic.replace(' ', '_')}_enhanced.pdf",
                        mime="application/pdf"
                    )
                    
                    st.success("✅ Enhanced PDF ready for download!")
                    
                except Exception as e:
                    st.error(f"Error generating PDF: {str(e)}")

if __name__ == "__main__":
    main()